"""
ATS-friendly Word documents: one column, standard fonts and headings, real
bullets, no tables, text boxes, images, headers or footers.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from jobpilot.tailoring.tailor import TailoredApplication

FONT = "Calibri"
BODY_SIZE = Pt(10.5)
TEXT_WIDTH = Inches(7.1)  # 8.5in page - 2 x 0.7in margins


def _base_document() -> Document:
    document = Document()
    for section in document.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(0.7)
        section.top_margin = section.bottom_margin = Inches(0.6)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = BODY_SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.08

    bullet = document.styles["List Bullet"]
    bullet.font.name = FONT
    bullet.font.size = BODY_SIZE
    bullet.paragraph_format.space_after = Pt(1)
    return document


def _bottom_border(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "808080")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    properties.append(borders)


def _heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11.5)
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    _bottom_border(paragraph)


def _left_right(document: Document, left: str, right: str, bold_left: bool = True) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    left_run = paragraph.add_run(left)
    left_run.bold = bold_left
    if right:
        paragraph.add_run("\t" + right)


def _bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item.lstrip("-•* ").strip(), style="List Bullet")


def _name_block(document: Document, application: TailoredApplication, centered: bool) -> None:
    name = document.add_paragraph()
    name.paragraph_format.space_after = Pt(0)
    run = name.add_run(application.candidate.name or "Your Name")
    run.bold = True
    run.font.size = Pt(18)
    if application.headline:
        headline = document.add_paragraph()
        headline.paragraph_format.space_after = Pt(0)
        headline_run = headline.add_run(application.headline)
        headline_run.font.size = Pt(11)
        if centered:
            headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if application.candidate.contact_line:
        contact = document.add_paragraph(application.candidate.contact_line)
        if centered:
            contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if centered:
        name.alignment = WD_ALIGN_PARAGRAPH.CENTER


WORD_PROPERTY_MAX = 255  # Word rejects longer core properties (title, keywords)


def fit_keywords(keywords: list[str], limit: int = WORD_PROPERTY_MAX) -> list[str]:
    """ As many keywords (max 25) as fit in ', '.join(...) <= limit."""
    kept: list[str] = []
    for word in keywords[:25]:
        if len(", ".join(kept + [word])) > limit:
            break
        kept.append(word)
    return kept


def _set_properties(document: Document, title: str, author: str, keywords: list[str]) -> None:
    if len(title) > WORD_PROPERTY_MAX:
        title = title[:WORD_PROPERTY_MAX - 3].rstrip() + "..."
    properties = document.core_properties
    properties.title = title
    properties.author = author
    properties.keywords = ", ".join(fit_keywords(list(keywords)))


def write_resume(application: TailoredApplication, path: Path, job_title: str) -> Path:
    document = _base_document()
    _name_block(document, application, centered=True)

    if application.summary:
        _heading(document, "Professional Summary")
        document.add_paragraph(application.summary)

    if application.skills:
        _heading(document, "Core Skills")
        document.add_paragraph(" | ".join(application.skills))

    if application.experience:
        _heading(document, "Professional Experience")
        for entry in application.experience:
            _left_right(document, entry.title, entry.dates)
            company = ", ".join(part for part in (entry.company, entry.location) if part)
            if company:
                line = document.add_paragraph()
                line.paragraph_format.space_after = Pt(1)
                line.add_run(company).italic = True
            _bullets(document, entry.bullets)

    if application.projects:
        _heading(document, "Projects")
        for project in application.projects:
            if project.name:
                _left_right(document, project.name, "")
            _bullets(document, project.bullets)

    if application.education:
        _heading(document, "Education")
        for school in application.education:
            _left_right(document, school.degree, school.dates)
            details = ", ".join(part for part in (school.institution, school.details) if part)
            if details:
                document.add_paragraph(details)

    if application.certifications:
        _heading(document, "Certifications")
        _bullets(document, application.certifications)

    _set_properties(
        document,
        f"{application.candidate.name} - Resume - {job_title}",
        application.candidate.name,
        application.matched_keywords,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def write_cover_letter(
    application: TailoredApplication,
    path: Path,
    company: str,
    job_title: str,
    today: date | None = None,
) -> Path:
    document = _base_document()
    for section in document.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(0.9)
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.space_after = Pt(8)

    _name_block(document, application, centered=False)
    document.add_paragraph((today or date.today()).strftime("%d %B %Y"))

    recipient = document.add_paragraph()
    recipient.add_run("Hiring Team")
    if company:
        recipient.add_run("\n" + company)
    subject = document.add_paragraph()
    subject.add_run(f"Re: Application for {job_title}").bold = True

    letter = application.cover_letter
    document.add_paragraph(letter.salutation or "Dear Hiring Manager,")
    for paragraph in letter.paragraphs:
        document.add_paragraph(paragraph)
    closing = document.add_paragraph(letter.closing or "Sincerely,")
    closing.paragraph_format.space_after = Pt(0)
    document.add_paragraph(application.candidate.name)

    _set_properties(
        document,
        f"{application.candidate.name} - Cover Letter - {job_title}",
        application.candidate.name,
        [],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path
