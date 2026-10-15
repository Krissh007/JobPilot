from __future__ import annotations

import re
from pathlib import Path

SUPPORTED_SUFFIXES = (".docx", ".pdf", ".txt", ".md")


class MasterResumeError(ValueError):
    """Raised when the master resume is missing or unreadable."""


def load_master_resume(path: str | Path) -> str:
    """Return the plain text of the original resume (.docx, .pdf, .txt, .md)."""
    resume_path = Path(path)
    if not resume_path.is_file():
        raise MasterResumeError(
            f"Master resume not found: {resume_path}. Put your original resume "
            "there or change tailoring.master_resume_path in config/settings.yaml."
        )

    suffix = resume_path.suffix.lower()
    if suffix == ".docx":
        text = _read_docx(resume_path)
    elif suffix == ".pdf":
        text = _read_pdf(resume_path)
    elif suffix in (".txt", ".md"):
        text = resume_path.read_text(encoding="utf-8")
    else:
        raise MasterResumeError(
            f"Unsupported resume format {suffix!r}; use one of {SUPPORTED_SUFFIXES}"
        )

    text = _clean(text)
    if len(text) < 200:
        raise MasterResumeError(
            f"Could not read enough text from {resume_path.name} "
            f"({len(text)} characters). If it is a scanned PDF, save it as .docx."
        )
    return text


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    lines = [paragraph.text for paragraph in document.paragraphs]
    # Many resume templates keep content in tables; read those too.
    for table in document.tables:
        for row in table.rows:
            seen: set[int] = set()
            for cell in row.cells:
                if id(cell._tc) in seen:  # merged cells repeat
                    continue
                seen.add(id(cell._tc))
                lines.append(cell.text)
    return "\n".join(lines)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _clean(text: str) -> str:
    text = text.replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()
