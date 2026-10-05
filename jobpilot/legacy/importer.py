"""
Import jobs the legacy agent found: its SQLite DB (opened read-only), or
the newest Excel export as a fallback. Status is mapped on first insert
only; after that JobPilot owns a job's stage.
"""

from __future__ import annotations

import re
import sqlite3
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable

from sqlmodel import Session

from jobpilot.db.models import Actor, JobStatus
from jobpilot.db.repo import upsert_job

IST = timezone(timedelta(hours=5, minutes=30))  # no DST, so a fixed offset is exact

# Legacy status -> JobPilot stage. None = do not import.
STATUS_MAP: dict[str, JobStatus | None] = {
    "new": JobStatus.found,
    "review": JobStatus.found,
    "shortlisted": JobStatus.shortlisted,
    "applied": JobStatus.applied,
    "interviewing": JobStatus.interview,
    "offer": JobStatus.offer,
    "skipped": JobStatus.rejected,  # your own decision in the legacy tools
    "rejected": None,  # dropped by the legacy quality gates
}
EXPERIENCE = re.compile(r"(\d+)\s*(?:-|to|–)\s*(\d+)")


@dataclass
class ImportStats:
    source: str = ""
    seen: int = 0
    created: int = 0
    updated: int = 0
    skipped: int = 0
    skipped_reasons: dict[str, int] = field(default_factory=dict)

    def skip(self, reason: str) -> None:
        self.skipped += 1
        self.skipped_reasons[reason] = self.skipped_reasons.get(reason, 0) + 1

    def summary(self) -> str:
        reasons = ", ".join(f"{k}: {v}" for k, v in self.skipped_reasons.items()) or "none"
        return (f"Imported from {self.source}: {self.seen} rows -> {self.created} new, "
                f"{self.updated} updated, {self.skipped} skipped ({reasons})")


def legacy_time(value: Any) -> datetime | None:
    """Legacy timestamps are naive local (IST) ISO strings."""
    if isinstance(value, datetime):
        parsed = value
    elif value:
        try:
            parsed = datetime.fromisoformat(str(value).strip())
        except ValueError:
            return None
    else:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=IST)
    return parsed.astimezone(timezone.utc)


def _int(value: Any) -> int | None:
    try:
        return int(float(value)) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _float(value: Any) -> float | None:
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def map_row(row: dict[str, Any]) -> tuple[dict[str, Any] | None, str]:
    """Legacy row (DB or Excel) -> JobPilot job dict, or (None, skip reason)."""
    url = str(row.get("job_url") or "").strip()
    if not url:
        return None, "no url"
    legacy_status = str(row.get("status") or "new").strip().lower()
    if legacy_status not in STATUS_MAP:
        legacy_status = "new"
    status = STATUS_MAP[legacy_status]
    if status is None:
        return None, f"legacy {legacy_status}"

    exp_min, exp_max = _int(row.get("experience_min")), _int(row.get("experience_max"))
    if exp_min is None and exp_max is None:
        match = EXPERIENCE.search(str(row.get("experience") or ""))
        if match:
            exp_min, exp_max = int(match.group(1)), int(match.group(2))

    job = {
        "url": url,
        "title": str(row.get("title") or "").strip() or "(untitled)",
        "company": str(row.get("company") or "").strip() or "(unknown)",
        "source": str(row.get("source") or "").strip().lower(),
        "location": str(row.get("location") or "").strip(),
        "exp_min": exp_min,
        "exp_max": exp_max,
        "jd_text": str(row.get("description") or "").strip(),
        "legacy_score": _float(row.get("match_score")),
        "status": status,
        "notes": str(row.get("notes") or "").strip(),
    }
    found_at = legacy_time(row.get("discovered_at"))
    if found_at:
        job["found_at"] = found_at
    if status in (JobStatus.applied, JobStatus.interview, JobStatus.offer):
        job["applied_at"] = legacy_time(row.get("applied_at")) or found_at
    return job, ""


def read_legacy_db(db_path: Path) -> list[dict[str, Any]]:
    if not db_path.is_file():
        raise FileNotFoundError(db_path)
    # Read-only URI: JobPilot can never write to the legacy database.
    uri = f"file:{db_path.as_posix()}?mode=ro"
    with sqlite3.connect(uri, uri=True, timeout=10) as connection:
        connection.row_factory = sqlite3.Row
        return [dict(row) for row in connection.execute("SELECT * FROM jobs ORDER BY id")]


def latest_excel(export_dir: Path) -> Path | None:
    files = [p for p in export_dir.glob("new_leads_*d_*.xlsx") if not p.name.startswith("~$")]
    return max(files, key=lambda p: p.stat().st_mtime, default=None)


def read_excel(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(rows, [])]
        return [dict(zip(header, values)) for values in rows if any(values)]
    finally:
        workbook.close()


def import_rows(session: Session, rows: Iterable[dict[str, Any]], source: str) -> ImportStats:
    stats = ImportStats(source=source)
    for row in rows:
        stats.seen += 1
        job, reason = map_row(row)
        if job is None:
            stats.skip(reason)
            continue
        _, created = upsert_job(session, job, actor=Actor.agent)
        if created:
            stats.created += 1
        else:
            stats.updated += 1
    return stats


def import_legacy(session: Session, legacy_dir: Path, log=print) -> ImportStats:
    """Prefer the legacy DB; fall back to the newest Excel export."""
    db_path = legacy_dir / "data" / "jobs.db"
    try:
        rows = read_legacy_db(db_path)
        source = str(db_path)
    except (OSError, sqlite3.Error) as error:
        log(f"Legacy DB unavailable ({error}); falling back to the newest Excel export")
        excel = latest_excel(legacy_dir / "data" / "exports")
        if excel is None:
            raise FileNotFoundError("No legacy jobs.db and no new_leads_*.xlsx export found") from error
        rows = read_excel(excel)
        source = str(excel)
    return import_rows(session, rows, source)
